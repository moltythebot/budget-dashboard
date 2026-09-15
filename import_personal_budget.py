"""Import the corrected Infinity Drones personal budget from the workbook.

This implements the rules recovered from the earlier correct review:
- Income = actual salary deposits in the cheque-credit account only
  (`FNB OB Pmt Salary ...` in the Cheque Credit columns), not the Salary sheet.
- Personal spending = cheque debits + credit-card debits after excluding business
  rows, SARS, credit-card clearing transfers, Pinegrove, Coenie Betaaling, Huis
  Betaal, investment-profit transfers, Tesla share, stock/courier/DJI/FPV rows,
  and maintenance above R4,500.
- The R1,400 "Scheduled Trf To Monthly 1K Transfer" is recorded as savings, not
  spending and not ignored.
- PSG AML R8,500 debit order is kept in spending as "Retirement (PSG Aml)" per
  Gerhard's prior instruction: it is a real monthly cash outflow until retirement.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
import sqlite3

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "budget.db"
EXCEL_PATH = Path("/home/molty/.hermes/cache/documents/doc_81dac821c761_2025_March_-_2026_Feb-_Provisional_tax_info_sheet_-_SR.xlsx")

CAT_FIX = {
    "insurnance": "insurance",
    "insurnace": "insurance",
    "maintennace": "maintenance",
    "accountng fees": "accounting fees",
}

DISPLAY_CATEGORY = {
    "accounting fees": "accounting fees",
    "bank fees": "bank fees",
    "cleaning": "cleaning",
    "donation": "donation",
    "electricity": "electricity",
    "entertainment": "entertainment",
    "insurance": "insurance",
    "internet": "internet",
    "maintenance": "maintenance",
    "meals": "meals",
    "medical": "medical",
    "rates": "rates",
    "security": "security",
    "subscription": "subscription",
    "telephone": "telephone",
    "travel": "travel",
    "uif": "uif",
}

BUSINESS_CATEGORIES = {"stock", "courier"}
BUSINESS_KEYWORDS = ["stock", "courier", "fpv", "dji", "air3s", "aliexpress", "bob group"]
TRANSFER_KEYWORDS = ["transfer", "scheduled trf", "scheduled transfer", "internal transfer", "monthly transfer"]


def clean_text(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def money(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    return None


def normalize_budget_date(dt: datetime) -> datetime:
    """Normalize source dates to the workbook period March 2025-February 2026.

    Some imported rows carried Aug-Dec as 2026 and one Feb row as 2025,
    which created empty visible months for Aug-Dec 2025. Keep the transaction
    values unchanged and only correct the fiscal-period year.
    """
    if dt.year == 2026 and 8 <= dt.month <= 12:
        return dt.replace(year=2025)
    if dt.year == 2025 and dt.month == 2:
        return dt.replace(year=2026)
    return dt


def month_key(dt: datetime) -> str:
    return normalize_budget_date(dt).strftime("%Y-%m")


def parse_cheque_credits(ws) -> list[dict]:
    rows = []
    for r in range(2, ws.max_row + 1):
        dt = ws.cell(r, 1).value
        desc = clean_text(ws.cell(r, 2).value)
        amount = money(ws.cell(r, 3).value)
        if amount is None:
            amount = money(ws.cell(r, 4).value)
        if isinstance(dt, datetime) and desc and amount:
            rows.append({"date": dt, "description": desc, "amount": amount, "account": "cheque_credit"})
    return rows


def parse_cheque_debits(ws) -> list[dict]:
    rows = []
    for r in range(2, ws.max_row + 1):
        dt = ws.cell(r, 6).value
        raw_desc = ws.cell(r, 7).value
        desc = clean_text(raw_desc)
        amount_h = money(ws.cell(r, 8).value)
        val_i = ws.cell(r, 9).value
        if not isinstance(dt, datetime) or raw_desc is None:
            continue

        if amount_h is not None:
            amount = amount_h
            category = clean_text(val_i).lower()
        else:
            amount_i = money(val_i)
            if amount_i is None:
                continue
            amount = amount_i
            category = ""

        category = CAT_FIX.get(category, category)
        if not desc and not category:
            continue
        rows.append({"date": dt, "description": desc, "amount": amount, "category": category, "account": "cheque_debit"})
    return rows


def parse_credit_card_debits(ws) -> list[dict]:
    rows = []
    for r in range(2, ws.max_row + 1):
        dt = ws.cell(r, 14).value
        desc = clean_text(ws.cell(r, 15).value)
        amount = money(ws.cell(r, 16).value)
        category = clean_text(ws.cell(r, 17).value).lower()
        if isinstance(dt, datetime) and desc and amount is not None:
            category = CAT_FIX.get(category, category)
            rows.append({"date": dt, "description": desc, "amount": amount, "category": category, "account": "credit_card_debit"})
    return rows


def is_salary_deposit(row: dict) -> bool:
    desc = row["description"].lower()
    return "salary" in desc and "ob pmt" in desc


def classify_outflow(row: dict) -> tuple[str, str, str, str] | None:
    desc = row["description"]
    dl = desc.lower()
    category = row.get("category", "")

    if "scheduled trf to monthly 1k transfer" in dl:
        return (
            "savings",
            "Monthly Transfer Savings",
            desc,
            "R1,400 monthly scheduled transfer; recorded as savings per Gerhard's correction.",
        )

    if category in BUSINESS_CATEGORIES:
        return None
    if any(keyword in dl for keyword in BUSINESS_KEYWORDS):
        return None
    if any(keyword in dl for keyword in TRANSFER_KEYWORDS):
        return None
    if "tesla share" in dl:
        return None
    if "investment" in dl and "profit" in dl:
        return None
    if "pinegrove" in dl:
        return None
    if "coenie betaaling" in dl:
        return None
    if "huis betaal" in dl:
        return None
    if category == "maintenance" and row["amount"] > 4500:
        return None
    if category == "sars":
        return None

    if "psg aml" in dl:
        return (
            "expense",
            "Retirement (PSG Aml)",
            desc,
            "Retirement annuity debit order; kept as spending/outflow per Gerhard's instruction.",
        )

    display = DISPLAY_CATEGORY.get(category, category or "Uncategorised")
    return (
        "expense",
        display,
        desc,
        "Imported from workbook after business/SARS/transfer exclusions.",
    )


def build_entries(wb) -> list[tuple]:
    ws = wb["March-Jan 2026"]
    cheque_credits = parse_cheque_credits(ws)
    outflows = parse_cheque_debits(ws) + parse_credit_card_debits(ws)

    entries: list[tuple] = []

    for row in cheque_credits:
        if is_salary_deposit(row):
            entries.append((
                month_key(row["date"]),
                "Cheque Credit / Net Salary",
                row["description"],
                row["amount"],
                "income",
                "Actual salary paid into cheque credit account; business income ignored.",
                "user_file",
                0,
            ))

    for row in outflows:
        if row["amount"] <= 0:
            continue
        classified = classify_outflow(row)
        if classified is None:
            continue
        entry_type, category, subcategory, notes = classified
        entries.append((
            month_key(row["date"]),
            category,
            subcategory,
            row["amount"],
            entry_type,
            notes,
            "user_file",
            0,
        ))

    return entries


def reset_imported_rows(conn: sqlite3.Connection):
    conn.execute("DELETE FROM budget_entries WHERE source IN ('seeded_sample', 'user_file')")
    conn.execute("DELETE FROM monthly_snapshots")


def insert_entries(conn: sqlite3.Connection, entries: list[tuple]):
    conn.executemany(
        """
        INSERT INTO budget_entries
            (month, category, subcategory, amount, entry_type, notes, source, is_sample)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        entries,
    )


def rebuild_snapshots(conn: sqlite3.Connection):
    rows = conn.execute(
        """
        SELECT month,
               SUM(CASE WHEN entry_type='income' THEN amount ELSE 0 END) AS income,
               SUM(CASE WHEN entry_type='expense' THEN amount ELSE 0 END) AS expenses,
               SUM(CASE WHEN entry_type='savings' THEN amount ELSE 0 END) AS savings
        FROM budget_entries
        GROUP BY month
        ORDER BY month
        """
    ).fetchall()

    for month, income, expenses, savings in rows:
        income = income or 0
        expenses = expenses or 0
        savings = savings or 0
        conn.execute(
            """
            INSERT INTO monthly_snapshots
                (month, total_income, total_expenses, total_savings, net_cashflow)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(month) DO UPDATE SET
                total_income=excluded.total_income,
                total_expenses=excluded.total_expenses,
                total_savings=excluded.total_savings,
                net_cashflow=excluded.net_cashflow
            """,
            (month, income, expenses, savings, income - expenses - savings),
        )


def print_verification(conn: sqlite3.Connection):
    print("Imported rows:", conn.execute("SELECT COUNT(*) FROM budget_entries WHERE source='user_file'").fetchone()[0])
    print("Sample rows:", conn.execute("SELECT COUNT(*) FROM budget_entries WHERE is_sample=1").fetchone()[0])
    print("By type:")
    for row in conn.execute("SELECT entry_type, COUNT(*), ROUND(SUM(amount), 2) FROM budget_entries GROUP BY entry_type ORDER BY entry_type"):
        print(f"  {row[0]}: rows={row[1]}, total=R{row[2]:,.2f}")

    print("\nMonthly summary:")
    rows = conn.execute(
        """
        SELECT month,
               ROUND(SUM(CASE WHEN entry_type='income' THEN amount ELSE 0 END), 2) AS income,
               ROUND(SUM(CASE WHEN entry_type='expense' THEN amount ELSE 0 END), 2) AS expenses,
               ROUND(SUM(CASE WHEN entry_type='savings' THEN amount ELSE 0 END), 2) AS savings,
               ROUND(SUM(CASE WHEN entry_type='income' THEN amount ELSE 0 END)
                 - SUM(CASE WHEN entry_type='expense' THEN amount ELSE 0 END)
                 - SUM(CASE WHEN entry_type='savings' THEN amount ELSE 0 END), 2) AS net
        FROM budget_entries
        GROUP BY month
        ORDER BY month
        """
    ).fetchall()
    for month, income, expenses, savings, net in rows:
        print(f"  {month}: income=R{income:,.2f}, expenses=R{expenses:,.2f}, savings=R{savings:,.2f}, net=R{net:,.2f}")

    totals = conn.execute(
        """
        SELECT ROUND(SUM(CASE WHEN entry_type='income' THEN amount ELSE 0 END), 2),
               ROUND(SUM(CASE WHEN entry_type='expense' THEN amount ELSE 0 END), 2),
               ROUND(SUM(CASE WHEN entry_type='savings' THEN amount ELSE 0 END), 2),
               ROUND(SUM(CASE WHEN entry_type='income' THEN amount ELSE 0 END)
                 - SUM(CASE WHEN entry_type='expense' THEN amount ELSE 0 END)
                 - SUM(CASE WHEN entry_type='savings' THEN amount ELSE 0 END), 2)
        FROM budget_entries
        """
    ).fetchone()
    print(f"\nTotals: income=R{totals[0]:,.2f}, expenses=R{totals[1]:,.2f}, savings=R{totals[2]:,.2f}, net=R{totals[3]:,.2f}")

    print("\nExpense categories:")
    for category, total, count in conn.execute("SELECT category, ROUND(SUM(amount), 2), COUNT(*) FROM budget_entries WHERE entry_type='expense' GROUP BY category ORDER BY SUM(amount) DESC"):
        print(f"  {category}: R{total:,.2f} ({count})")


def main():
    if not EXCEL_PATH.exists():
        raise SystemExit(f"Workbook not found: {EXCEL_PATH}")

    wb = load_workbook(EXCEL_PATH, data_only=True)
    entries = build_entries(wb)

    conn = sqlite3.connect(DB_PATH)
    try:
        reset_imported_rows(conn)
        insert_entries(conn, entries)
        rebuild_snapshots(conn)
        conn.commit()
        print_verification(conn)
    finally:
        conn.close()


if __name__ == "__main__":
    main()
